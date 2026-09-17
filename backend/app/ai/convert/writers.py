"""把 IR 渲染成目标格式（md / docx / pdf）。"""

import io
from xml.sax.saxutils import escape

from app.ai.convert.ir import (
    Block,
    CodeBlock,
    Heading,
    ListBlock,
    PageBreak,
    Paragraph,
    Quote,
    Rule,
    Span,
    TableBlock,
    spans_text,
)

# ---------- 公共 ----------


def _table_rows(block: TableBlock) -> list[list[str]]:
    """补齐每行列数到等宽，返回 [表头, *数据行]；空表返回 []。"""
    width = max([len(block.header), *(len(row) for row in block.rows)], default=0)
    if width == 0:
        return []
    header = list(block.header) + [""] * (width - len(block.header))
    rows = [list(row) + [""] * (width - len(row)) for row in block.rows]
    return [header, *rows]


# ---------- Markdown ----------


def _md_span(span: Span) -> str:
    if span.code:
        return f"`{span.text}`"
    if span.bold and span.italic:
        return f"***{span.text}***"
    if span.bold:
        return f"**{span.text}**"
    if span.italic:
        return f"*{span.text}*"
    return span.text


def _md_spans(spans: list[Span]) -> str:
    return "".join(_md_span(span) for span in spans)


def _md_cell(text: str) -> str:
    """仅转义 markdown 管道表的单元格：竖线会切断列、换行会切断行、反斜杠是转义符。
    只用于 md 渲染路径，docx/pdf 共用 _table_rows 的原文，不能在此处转义。"""
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\n", " ")


def to_markdown(blocks: list[Block]) -> bytes:
    lines: list[str] = []
    for block in blocks:
        if isinstance(block, Heading):
            lines += ["#" * min(max(block.level, 1), 6) + " " + _md_spans(block.spans), ""]
        elif isinstance(block, Paragraph):
            lines += [_md_spans(block.spans), ""]
        elif isinstance(block, ListBlock):
            for index, item in enumerate(block.items, 1):
                marker = f"{index}." if block.ordered else "-"
                lines.append(f"{marker} {_md_spans(item)}")
            lines.append("")
        elif isinstance(block, CodeBlock):
            fence = f"```{block.language}" if block.language else "```"
            lines += [fence, block.text, "```", ""]
        elif isinstance(block, Quote):
            lines += [f"> {_md_spans(block.spans)}", ""]
        elif isinstance(block, TableBlock):
            rows = _table_rows(block)
            if rows:
                lines.append("| " + " | ".join(_md_cell(cell) for cell in rows[0]) + " |")
                lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
                lines.extend(
                    "| " + " | ".join(_md_cell(cell) for cell in row) + " |" for row in rows[1:]
                )
                lines.append("")
        elif isinstance(block, (PageBreak, Rule)):
            lines += ["---", ""]
    text = "\n".join(lines).strip()
    return (text + "\n").encode("utf-8") if text else b"\n"


# ---------- DOCX ----------


def _fill_runs(paragraph, spans: list[Span]) -> None:
    for span in spans:
        run = paragraph.add_run(span.text)
        run.bold = span.bold
        run.italic = span.italic
        if span.code:
            run.font.name = "Consolas"


def to_docx(blocks: list[Block]) -> bytes:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    for block in blocks:
        if isinstance(block, Heading):
            document.add_heading(spans_text(block.spans), level=min(max(block.level, 1), 6))
        elif isinstance(block, Paragraph):
            _fill_runs(document.add_paragraph(), block.spans)
        elif isinstance(block, ListBlock):
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                _fill_runs(document.add_paragraph(style=style), item)
        elif isinstance(block, CodeBlock):
            run = document.add_paragraph().add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif isinstance(block, Quote):
            _fill_runs(document.add_paragraph(style="Quote"), block.spans)
        elif isinstance(block, TableBlock):
            rows = _table_rows(block)
            if rows:
                table = document.add_table(rows=1, cols=len(rows[0]))
                table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, rows[0], strict=False):
                    cell.text = text
                for row in rows[1:]:
                    for cell, text in zip(table.add_row().cells, row, strict=False):
                        cell.text = text
        elif isinstance(block, PageBreak):
            document.add_page_break()
        elif isinstance(block, Rule):
            document.add_paragraph("―" * 10)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------- PDF ----------

_CJK_FONT = "STSong-Light"
# reportlab 的字体注册是进程级全局状态：只注册一次、幂等；并发首次调用可能重复注册
# （相同字体名 + 相同定义，reportlab 幂等处理），无副作用。
_cjk_registered = False


def _ensure_cjk_font() -> None:
    """注册 reportlab 内置简体中文 CID 字体，无需外部字体文件。"""
    global _cjk_registered
    if _cjk_registered:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    pdfmetrics.registerFont(UnicodeCIDFont(_CJK_FONT))
    _cjk_registered = True


def to_pdf(blocks: list[Block]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        SimpleDocTemplate,
        Spacer,
        TableStyle,
    )
    from reportlab.platypus import PageBreak as RLPageBreak
    from reportlab.platypus import Paragraph as RLParagraph
    from reportlab.platypus import Table as RLTable

    _ensure_cjk_font()
    base = getSampleStyleSheet()["BodyText"]
    body = ParagraphStyle("cjk-body", parent=base, fontName=_CJK_FONT, fontSize=10.5, leading=16)
    quote = ParagraphStyle(
        "cjk-quote", parent=body, leftIndent=12, textColor=colors.HexColor("#555555")
    )
    code = ParagraphStyle("cjk-code", parent=body, fontName="Courier", fontSize=9, leading=12)
    heading_sizes = {1: 20, 2: 16, 3: 14, 4: 12, 5: 11, 6: 11}

    flow: list = []
    for block in blocks:
        if isinstance(block, Heading):
            level = min(max(block.level, 1), 6)
            style = ParagraphStyle(
                f"cjk-h{level}",
                parent=body,
                fontName=_CJK_FONT,
                fontSize=heading_sizes[level],
                leading=heading_sizes[level] + 6,
                spaceBefore=10,
                spaceAfter=6,
            )
            flow.append(RLParagraph(escape(spans_text(block.spans)), style))
        elif isinstance(block, Paragraph):
            flow.append(RLParagraph(escape(spans_text(block.spans)), body))
        elif isinstance(block, ListBlock):
            items = [ListItem(RLParagraph(escape(spans_text(item)), body)) for item in block.items]
            if items:
                flow.append(ListFlowable(items, bulletType="1" if block.ordered else "bullet"))
        elif isinstance(block, CodeBlock):
            for line in block.text.split("\n"):
                flow.append(RLParagraph(escape(line) or "&nbsp;", code))
            flow.append(Spacer(1, 4))
        elif isinstance(block, Quote):
            flow.append(RLParagraph(escape(spans_text(block.spans)), quote))
        elif isinstance(block, TableBlock):
            rows = _table_rows(block)
            if rows:
                data = [[RLParagraph(escape(cell), body) for cell in row] for row in rows]
                table = RLTable(data, repeatRows=1)
                table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    )
                )
                flow.append(table)
                flow.append(Spacer(1, 6))
        elif isinstance(block, (PageBreak, Rule)):
            if flow:
                flow.append(RLPageBreak())
    if not flow:
        flow.append(Spacer(1, 1))
    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    ).build(flow)
    return buffer.getvalue()
