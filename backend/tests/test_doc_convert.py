"""文档格式转换：pdf/docx/md/txt → md/docx/pdf（docs/设计/25）。"""

import io

import pytest

from app.ai.convert import ConvertError, convert_file

MD_SAMPLE = (
    "# 标题\n\n"
    "正文包含**粗体**与 `代码`。\n\n"
    "- 甲\n- 乙\n\n"
    "| 列1 | 列2 |\n| --- | --- |\n| a | b |\n\n"
    "```python\nprint(1)\n```\n"
)


def write(tmp_path, name: str, data: str | bytes):
    path = tmp_path / name
    if isinstance(data, str):
        path.write_text(data, encoding="utf-8")
    else:
        path.write_bytes(data)
    return path


def test_markdown_to_docx_keeps_structure(tmp_path):
    out = convert_file(write(tmp_path, "note.md", MD_SAMPLE), "md", "docx")
    assert out.ext == "docx"
    assert out.mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    import docx

    document = docx.Document(io.BytesIO(out.data))
    headings = [p for p in document.paragraphs if p.style.name == "Heading 1"]
    assert [p.text for p in headings] == ["标题"]
    assert any("正文包含" in p.text for p in document.paragraphs)
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "列1"


def test_txt_to_markdown_splits_paragraphs(tmp_path):
    out = convert_file(write(tmp_path, "a.txt", "第一段\n\n第二段"), "txt", "md")
    text = out.data.decode("utf-8")
    assert "第一段" in text and "第二段" in text


def test_same_format_and_unknown_target_rejected(tmp_path):
    src = write(tmp_path, "note.md", "# x")
    with pytest.raises(ConvertError, match="已经是 md"):
        convert_file(src, "md", "md")
    with pytest.raises(ConvertError, match="目标格式"):
        convert_file(src, "md", "html")


def test_markdown_to_docx_keeps_list_and_code(tmp_path):
    out = convert_file(write(tmp_path, "note.md", MD_SAMPLE), "md", "docx")
    import docx

    document = docx.Document(io.BytesIO(out.data))
    assert "List Bullet" in [p.style.name for p in document.paragraphs]
    assert any("print(1)" in p.text for p in document.paragraphs)


def test_txt_to_markdown_keeps_paragraph_break(tmp_path):
    out = convert_file(write(tmp_path, "a.txt", "第一段\n\n第二段"), "txt", "md")
    assert "第一段\n\n第二段" in out.data.decode("utf-8")


def test_trailing_line_break_does_not_add_blank_line():
    from app.ai.convert.readers import read_markdown
    from app.ai.convert.writers import to_markdown

    text = to_markdown(read_markdown("第一段<br>\n\n第二段")).decode("utf-8")
    assert "第一段\n\n\n" not in text
    assert "第一段\n\n第二段" in text


def test_txt_to_docx(tmp_path):
    out = convert_file(write(tmp_path, "a.txt", "第一段\n\n第二段"), "txt", "docx")
    import docx

    document = docx.Document(io.BytesIO(out.data))
    assert any("第一段" in p.text for p in document.paragraphs)


def test_docx_to_markdown_keeps_list_and_code(tmp_path):
    docx_bytes = convert_file(write(tmp_path, "note.md", MD_SAMPLE), "md", "docx").data
    text = convert_file(write(tmp_path, "note.docx", docx_bytes), "docx", "md").data.decode("utf-8")
    assert "print(1)" in text
    assert "- 甲" in text


def test_empty_content_is_rejected(tmp_path):
    with pytest.raises(ConvertError, match="未提取到"):
        convert_file(write(tmp_path, "empty.txt", "   \n\n"), "txt", "md")


def test_missing_file_and_unsupported_source(tmp_path):
    with pytest.raises(ConvertError, match="源文件不存在"):
        convert_file(tmp_path / "nope.md", "md", "docx")
    src = write(tmp_path, "a.rtf", "x")
    with pytest.raises(ConvertError, match="不支持的来源格式"):
        convert_file(src, "rtf", "md")


def test_nested_list_items_do_not_merge():
    from app.ai.convert.readers import read_markdown
    from app.ai.convert.writers import to_markdown

    text = to_markdown(read_markdown("- A\n    - A1\n- B\n")).decode("utf-8")
    assert "AA1" not in text
    assert "A1" in text


def test_markdown_table_cell_pipe_is_escaped():
    from app.ai.convert.ir import TableBlock
    from app.ai.convert.writers import to_markdown

    text = to_markdown([TableBlock(header=["a|b", "c"], rows=[])]).decode("utf-8")
    assert "a\\|b" in text


def test_docx_to_markdown_round_trip(tmp_path):
    docx_bytes = convert_file(write(tmp_path, "note.md", MD_SAMPLE), "md", "docx").data
    out = convert_file(write(tmp_path, "note.docx", docx_bytes), "docx", "md")
    text = out.data.decode("utf-8")
    assert "# 标题" in text
    assert "正文包含" in text
    assert "| 列1 | 列2 |" in text


def test_docx_to_html_contains_semantics(tmp_path):
    from app.ai.convert.readers import docx_to_html

    docx_bytes = convert_file(write(tmp_path, "note.md", "# 甲\n\n乙"), "md", "docx").data
    html = docx_to_html(write(tmp_path, "note.docx", docx_bytes))
    assert "<h1>" in html and "甲" in html


def make_pdf(tmp_path, text: str):
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    # 用默认 Helvetica 写 ASCII，保证抽取得回原文字（CID 字体的抽取不稳定，见 Task 4 说明）
    page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return write(tmp_path, "src.pdf", data)


def test_pdf_to_markdown(tmp_path):
    out = convert_file(make_pdf(tmp_path, "PDF body hello"), "pdf", "md")
    assert "PDF body hello" in out.data.decode("utf-8")


def test_pdf_to_docx(tmp_path):
    import docx

    out = convert_file(make_pdf(tmp_path, "PDF body hello"), "pdf", "docx")
    document = docx.Document(io.BytesIO(out.data))
    assert any("PDF body hello" in p.text for p in document.paragraphs)


def test_empty_pdf_is_rejected(tmp_path):
    import fitz

    doc = fitz.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    with pytest.raises(ConvertError, match="未提取到"):
        convert_file(write(tmp_path, "empty.pdf", data), "pdf", "md")


def test_markdown_to_pdf(tmp_path):
    out = convert_file(
        write(tmp_path, "n.md", "# Title\n\nBody text\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n"),
        "md",
        "pdf",
    )
    assert out.data[:5] == b"%PDF-"
    # CID 字体（STSong-Light）的文本抽取依赖阅读器内置 CMap，断言不可靠；
    # 只验证是合法 PDF 且页数正常，"内容不丢"由 md/docx 目标的用例覆盖。
    import fitz

    with fitz.open(stream=out.data, filetype="pdf") as pdf:
        assert pdf.page_count >= 1


def test_markdown_to_pdf_with_chinese(tmp_path):
    # 中文走 CID 内置字体：只验证不报错且是合法 PDF（CID 文本抽取不可靠，不断言原文）
    out = convert_file(write(tmp_path, "c.md", "# 中文标题\n\n正文内容"), "md", "pdf")
    assert out.data[:5] == b"%PDF-"


def test_docx_to_pdf(tmp_path):
    docx_bytes = convert_file(write(tmp_path, "n.md", "# 甲\n\n乙"), "md", "docx").data
    out = convert_file(write(tmp_path, "n.docx", docx_bytes), "docx", "pdf")
    assert out.data[:5] == b"%PDF-"


def test_markdown_to_pdf_long_table_cell_does_not_crash(tmp_path):
    cell = "很长的说明文字" * 20
    md = f"| 名称 | 说明 |\n| --- | --- |\n| A | {cell} |\n"
    out = convert_file(write(tmp_path, "t.md", md), "md", "pdf")
    assert out.data[:5] == b"%PDF-"


def test_pdf_rule_is_separator_not_page_break(tmp_path):
    out = convert_file(write(tmp_path, "r.md", "第一段\n\n---\n\n第二段"), "md", "pdf")
    import fitz

    with fitz.open(stream=out.data, filetype="pdf") as pdf:
        assert pdf.page_count == 1
