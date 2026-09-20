"""doc_create 工具：Markdown 正文直接渲染成文件（docs/设计/25 同一条渲染链）。"""

import io

import pytest

from app.ai.convert import ConvertError, convert_markdown

MD_SAMPLE = (
    "# 请假条\n\n"
    "尊敬的领导：\n\n"
    "本人因感冒发烧，特申请请假一天。\n\n"
    "- 请假人：张三\n- 日期：2026-09-21\n\n"
    "| 事项 | 内容 |\n| --- | --- |\n| 天数 | 1 天 |\n"
)


def test_markdown_to_docx_keeps_structure():
    out = convert_markdown(MD_SAMPLE, "docx")
    assert out.ext == "docx"
    assert out.mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    import docx

    document = docx.Document(io.BytesIO(out.data))
    headings = [p for p in document.paragraphs if p.style.name == "Heading 1"]
    assert [p.text for p in headings] == ["请假条"]
    assert any("张三" in p.text for p in document.paragraphs)
    assert len(document.tables) == 1


def test_markdown_to_pdf_is_valid_pdf():
    out = convert_markdown(MD_SAMPLE, "pdf")
    assert out.ext == "pdf"
    assert out.data[:5] == b"%PDF-"

    import fitz

    with fitz.open(stream=out.data, filetype="pdf") as pdf:
        assert pdf.page_count >= 1


def test_markdown_to_md_round_trip():
    out = convert_markdown(MD_SAMPLE, "md")
    assert out.ext == "md"
    text = out.data.decode("utf-8")
    assert "# 请假条" in text and "张三" in text


def test_empty_content_is_rejected():
    with pytest.raises(ConvertError, match="未提供"):
        convert_markdown("   \n\n", "docx")


def test_invalid_target_is_rejected():
    with pytest.raises(ConvertError, match="目标格式"):
        convert_markdown(MD_SAMPLE, "html")
